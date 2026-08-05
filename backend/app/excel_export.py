"""Generate the summary workbook (one sheet per financial year) for a category.

Layout mirrors the client's Hair-Oil sample exactly:

    CATEGORY | BRAND | COMMERCIAL | <year> months... | YTD

- No Mont/Weekly-Avg columns.
- Only the months that have data are shown (year start -> current month).
- Per brand: commercial themes, optional Tag / Value Adds rows, Total Spends,
  ACD (Com Only), ACD (All Exp); then TOTAL CATEGORY SPEND and SOS %.
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Green palette matching the client's sample.
FILL_HEADER = PatternFill("solid", fgColor="00B050")   # bright green header
FILL_TOTAL = PatternFill("solid", fgColor="E2EFC7")    # pale yellow-green Total row
FILL_ACD = PatternFill("solid", fgColor="C6EFCE")      # light green ACD rows
FILL_CATTOTAL = PatternFill("solid", fgColor="00B050")
FILL_SOS = PatternFill("solid", fgColor="C6EFCE")

FMT_INT = '_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'
FMT_PCT = "0%"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
VCENTER = Alignment(vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")


def build_workbook(summaries: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for summ in summaries:
        _write_sheet(wb, summ)
    if not wb.worksheets:
        wb.create_sheet("Summary")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _write_sheet(wb, summ):
    title = summ["fy"][:31]
    ws = wb.create_sheet(title=title)
    keys = summ["month_keys"]
    n = max(len(keys), 1)
    FIRST_M = 4              # column D
    ytd_col = FIRST_M + n    # column after the months

    # ---------- header (2 rows) ----------
    ws.cell(row=1, column=1, value="CATEGORY")
    ws.cell(row=1, column=2, value="BRAND")
    ws.cell(row=1, column=3, value="COMMERCIAL")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    ws.merge_cells(start_row=1, start_column=3, end_row=2, end_column=3)
    yr = ws.cell(row=1, column=FIRST_M, value=summ["fy"].replace("-", "/"))
    ws.merge_cells(start_row=1, start_column=FIRST_M, end_row=1, end_column=ytd_col - 1)
    ws.cell(row=1, column=ytd_col, value="YTD")
    ws.merge_cells(start_row=1, start_column=ytd_col, end_row=2, end_column=ytd_col)
    for i, mm in enumerate(summ["months"]):
        ws.cell(row=2, column=FIRST_M + i, value=mm["header"])
    for r in (1, 2):
        for c in range(1, ytd_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(bold=True, size=10)
            cell.fill = FILL_HEADER
            cell.alignment = CENTER
            cell.border = BORDER

    row = 3
    data_start = row

    def put_row(r, label_col, label, data, *, fill=None, bold=False, blank_zero=False,
                pct=False, show_ytd=True):
        cell = ws.cell(row=r, column=label_col, value=label)
        cell.font = Font(bold=bold, size=10)
        cell.alignment = LEFT
        if fill:
            cell.fill = fill
        for i, k in enumerate(keys):
            v = data.get(k)
            if blank_zero and (v == 0 or v is None):
                v = None
            cc = ws.cell(row=r, column=FIRST_M + i, value=v)
            cc.number_format = FMT_PCT if pct else FMT_INT
            cc.font = Font(bold=bold, size=10)
            if fill:
                cc.fill = fill
            cc.border = BORDER
        yv = data.get("ytd") if show_ytd else None
        if blank_zero and (yv == 0 or yv is None):
            yv = None
        yc = ws.cell(row=r, column=ytd_col, value=yv)
        yc.number_format = FMT_PCT if pct else FMT_INT
        yc.font = Font(bold=bold, size=10)
        if fill:
            yc.fill = fill
        yc.border = BORDER
        # borders on empty text columns for a clean grid
        for c in range(1, label_col):
            ws.cell(row=r, column=c).border = BORDER

    for g in summ["groups"]:
        for b in g["brands"]:
            brand_start = row
            # commercial themes with any spend in view
            themes = [t for t in b["themes"] if (t.get("ytd") or 0) != 0]
            ws.cell(row=brand_start, column=2, value=b["brand"]).font = Font(bold=True, size=10)
            for t in themes:
                put_row(row, 3, t["text"], t, blank_zero=True)
                row += 1
            if not themes:
                # keep an anchor row so the brand label has a cell
                put_row(row, 3, "", {}, blank_zero=True)
                row += 1
            if (b["tag"].get("ytd") or 0) != 0:
                put_row(row, 3, "Tag", b["tag"], blank_zero=True)
                row += 1
            if (b["value_adds"].get("ytd") or 0) != 0:
                put_row(row, 3, "Value Adds", b["value_adds"], blank_zero=True)
                row += 1
            put_row(row, 3, "Total Spends (000)", b["total"], fill=FILL_TOTAL, bold=True)
            row += 1
            put_row(row, 3, "ACD (Com Only)", b["acd_com"], fill=FILL_ACD, bold=True,
                    show_ytd=False)
            row += 1
            put_row(row, 3, "ACD (All Exp)", b["acd_all"], fill=FILL_ACD, bold=True,
                    show_ytd=False)
            row += 1
            # merge the brand label down its block
            ws.merge_cells(start_row=brand_start, start_column=2, end_row=row - 1, end_column=2)
            ws.cell(row=brand_start, column=2).alignment = CENTER

        if g["show_total"]:
            put_row(row, 2, f"Total {g['mother_brand']}", g["total"],
                    fill=FILL_TOTAL, bold=True)
            row += 1
            put_row(row, 2, f"Total {g['mother_brand']} - ACD", g["total_acd"],
                    fill=FILL_ACD, bold=True, show_ytd=False)
            row += 1

    data_end = row - 1
    # merge CATEGORY (product group) down the whole block
    if data_end >= data_start:
        ws.cell(row=data_start, column=1, value=summ["category"]).font = Font(bold=True, size=10)
        ws.merge_cells(start_row=data_start, start_column=1, end_row=data_end, end_column=1)
        ws.cell(row=data_start, column=1).alignment = CENTER

    # ---------- TOTAL CATEGORY SPEND ----------
    put_row(row, 1, "TOTAL CATEGORY SPEND", summ["category_total"],
            fill=FILL_CATTOTAL, bold=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    for c in range(1, ytd_col + 1):
        ws.cell(row=row, column=c).fill = FILL_CATTOTAL
        ws.cell(row=row, column=c).border = BORDER
    ws.cell(row=row, column=1).alignment = CENTER
    row += 1

    # ---------- SOS % ----------
    sos_start = row
    for s in summ["sos"]:
        ws.cell(row=row, column=2, value=s["mother_brand"]).font = Font(bold=True, size=10)
        ws.cell(row=row, column=2).alignment = LEFT
        ws.cell(row=row, column=2).fill = FILL_SOS
        put_row(row, 3, "", s, fill=FILL_SOS, pct=True)
        ws.cell(row=row, column=1).fill = FILL_SOS
        row += 1
    # merge "SOS %" label across the SOS rows (column C)
    if row - 1 >= sos_start:
        ws.cell(row=sos_start, column=3, value="SOS %").font = Font(bold=True, size=10)
        ws.merge_cells(start_row=sos_start, start_column=3, end_row=row - 1, end_column=3)
        ws.cell(row=sos_start, column=3).alignment = CENTER
    # final 100% total row
    total_pct = {k: (1.0 if any(s.get(k) for s in summ["sos"]) else None) for k in keys}
    total_pct["ytd"] = 1.0 if summ["sos"] else None
    put_row(row, 3, "", total_pct, fill=FILL_SOS, pct=True)
    for c in (1, 2):
        ws.cell(row=row, column=c).fill = FILL_SOS
        ws.cell(row=row, column=c).border = BORDER
    row += 1

    # ---------- widths / freeze ----------
    widths = {1: 15, 2: 24, 3: 60}
    for i in range(FIRST_M, ytd_col + 1):
        widths[i] = 11
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"{get_column_letter(FIRST_M)}3"
