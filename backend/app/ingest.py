"""Parse an uploaded TV / Radio / Press export and aggregate it.

Media type is auto-detected from the columns present:
    - has a "Channel" column  -> Radio   (no per-spot Brand; brand = mother brand)
    - has a "Dur(secs)" column -> TV
    - has an "Ins" column      -> Press   (insertions, no duration/freq)
"""
import io
import re
from collections import defaultdict
from datetime import date, datetime

import openpyxl

MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def _norm(s) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def _num(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("(", "-").replace(")", "").strip()
    if s in ("", "-", "."):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _parse_month(v) -> str | None:
    """Return 'YYYY-MM' from many month representations.

    Handles real Excel dates (the common case: a 'Jun-26'-formatted cell is
    actually the date 2026-06-01), plus text like 'Jun-26', 'July 2026',
    '2026-06', '06/2026'.
    """
    if v is None:
        return None
    # Real date / datetime cell (openpyxl returns these for mmm-yy formatted cells).
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v).strip()
    if not s:
        return None
    # Month name + year, e.g. "Jun-26", "July 2026", "Jun/26".
    m = re.match(r"([A-Za-z]{3,})[\s\-/.]+(\d{2,4})", s)
    if m:
        mon = MONTHS.get(m.group(1)[:3].lower())
        if mon:
            yr = int(m.group(2))
            return f"{(yr + 2000) if yr < 100 else yr:04d}-{mon:02d}"
    # Numeric year-month, e.g. "2026-06", "2026/06".
    m = re.match(r"(\d{4})[\s\-/.](\d{1,2})", s)
    if m:
        yr, mon = int(m.group(1)), int(m.group(2))
        if 1 <= mon <= 12:
            return f"{yr:04d}-{mon:02d}"
    # Numeric month-year, e.g. "06/2026", "06-26".
    m = re.match(r"(\d{1,2})[\s\-/.](\d{2,4})", s)
    if m:
        mon, yr = int(m.group(1)), int(m.group(2))
        if 1 <= mon <= 12:
            return f"{(yr + 2000) if yr < 100 else yr:04d}-{mon:02d}"
    return None


# Map many possible header spellings to canonical field names.
HEADER_ALIASES = {
    "month": "month",
    "product group": "product_group",
    "productgroup": "product_group",
    "category": "product_group",
    "mother brand": "mother_brand",
    "motherbrand": "mother_brand",
    "advertiser": "advertiser",
    "brand": "brand",
    "theme": "theme",
    "channel": "channel",
    "publication": "publication",
    "(000rs)": "spend",
    "000rs": "spend",
    "spend": "spend",
    "frq": "freq",
    "freq": "freq",
    "dur(secs)": "duration",
    "dur (secs)": "duration",
    "duration": "duration",
    "ins": "insertions",
    "insertions": "insertions",
}


_MEDIA_HEADERS = {"channel", "dur(secs)", "dur (secs)", "duration", "ins",
                  "insertions", "(000rs)", "000rs"}


def _find_header_row(ws):
    """Locate the header row: it contains 'Month' plus either a category column
    or a recognisable media/spend column."""
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = [_norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]
        if "month" not in vals:
            continue
        if any(v in ("product group", "category") for v in vals) or any(
            v in _MEDIA_HEADERS for v in vals
        ):
            return r, vals
    return None, None


def parse_workbook(content: bytes):
    """Return (media_type, aggregates, months, categories, row_count, warnings).

    `aggregates` is a dict keyed by
        (category, mother_brand, brand, theme_raw, month)
    -> {spend, freq, duration, insertions}
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    hdr_row, hdr_vals = _find_header_row(ws)
    if hdr_row is None:
        raise ValueError(
            "Could not find a header row containing 'Month' and 'Product Group'."
        )

    col_map = {}  # field -> column index (1-based)
    for idx, name in enumerate(hdr_vals, start=1):
        field = HEADER_ALIASES.get(name)
        if field and field not in col_map:
            col_map[field] = idx

    if "month" not in col_map:
        raise ValueError("No 'Month' column found.")

    # Detect media type.
    if "channel" in col_map:
        media_type = "radio"
    elif "insertions" in col_map:
        media_type = "press"
    elif "duration" in col_map:
        media_type = "tv"
    else:
        raise ValueError(
            "Could not detect media type. Expected a Channel (radio), "
            "Dur(secs) (TV) or Ins (press) column."
        )

    aggs = defaultdict(lambda: {"spend": 0.0, "freq": 0, "duration": 0, "insertions": 0})
    months, categories = set(), set()
    warnings = []
    row_count = 0
    skipped_no_month = 0
    skipped_no_mother = 0
    sample_months = []  # raw Month values we failed to parse, for diagnostics

    def cell(row, field):
        c = col_map.get(field)
        return ws.cell(row=row, column=c).value if c else None

    for r in range(hdr_row + 1, ws.max_row + 1):
        raw_month = cell(r, "month")
        month = _parse_month(raw_month)
        if not month:
            # Only count rows that actually have some content (skip blank rows).
            if raw_month not in (None, "") or cell(r, "mother_brand") or cell(r, "spend"):
                skipped_no_month += 1
                if raw_month is not None and len(sample_months) < 5:
                    sample_months.append(repr(raw_month))
            continue
        category = (str(cell(r, "product_group") or "").strip()) or "Uncategorised"
        mother = (str(cell(r, "mother_brand") or "").strip()) or (
            str(cell(r, "advertiser") or "").strip()
        )
        if not mother:
            skipped_no_mother += 1
            continue
        # Radio feed has no per-spot Brand -> brand defaults to the mother brand.
        brand = (str(cell(r, "brand") or "").strip()) or mother
        theme = str(cell(r, "theme") or "").strip()
        if not theme:
            # Press feeds have no Theme column; use the Publication as the label.
            theme = str(cell(r, "publication") or "").strip() or "(no theme)"

        key = (category, mother, brand, theme, month)
        a = aggs[key]
        a["spend"] += _num(cell(r, "spend"))
        a["freq"] += int(_num(cell(r, "freq")))
        a["duration"] += int(_num(cell(r, "duration")))
        a["insertions"] += int(_num(cell(r, "insertions")))

        months.add(month)
        categories.add(category)
        row_count += 1

    wb.close()
    if row_count == 0:
        detail = (
            f"No data rows were parsed. Detected columns: {sorted(col_map.keys())}. "
            f"Rows skipped because the Month couldn't be read: {skipped_no_month}"
        )
        if sample_months:
            detail += f" (sample Month values: {', '.join(sample_months)})"
        if skipped_no_mother:
            detail += f"; rows skipped for missing Mother Brand/Advertiser: {skipped_no_mother}"
        warnings.append(detail)
    elif skipped_no_month:
        warnings.append(
            f"{skipped_no_month} row(s) were skipped because their Month could not be read."
        )
    return media_type, aggs, sorted(months), sorted(categories), row_count, warnings
